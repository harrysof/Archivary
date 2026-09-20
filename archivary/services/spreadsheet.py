"""CSV / Excel import and export helpers.

Used by the batch-upload and batch-metadata features (mirroring
``ia upload --spreadsheet``) and by the search-results export button.
"""

from __future__ import annotations

import csv
import json
import logging
from collections.abc import Iterable, Sequence
from pathlib import Path

log = logging.getLogger(__name__)

IDENTIFIER_COLUMNS = ("identifier", "id", "item", "item_identifier")
FILE_COLUMNS = ("file", "files", "path", "paths", "filename", "filenames")
LIST_SEPARATOR = "|"

# Columns that are interpreted as file selections or the item key rather than
# as metadata to send to the Internet Archive.
RESERVED_COLUMNS = set(IDENTIFIER_COLUMNS) | set(FILE_COLUMNS) | {"_files", "_identifier"}


def read_table(path: str) -> tuple[list[str], list[dict[str, str]]]:
    """Read a CSV or Excel file into a header list and row dictionaries."""
    file_path = Path(path)
    suffix = file_path.suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        return _read_xlsx(file_path)
    return _read_csv(file_path)


def _read_csv(path: Path) -> tuple[list[str], list[dict[str, str]]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        sample = handle.read(4096)
        handle.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        reader = csv.DictReader(handle, dialect=dialect)
        headers = [h.strip() for h in (reader.fieldnames or []) if h is not None]
        rows: list[dict[str, str]] = []
        for raw in reader:
            row = {
                (key.strip() if key else ""): (value or "").strip()
                for key, value in raw.items()
                if key is not None
            }
            if any(v for v in row.values()):
                rows.append(row)
    return headers, rows


def _read_xlsx(path: Path) -> tuple[list[str], list[dict[str, str]]]:
    try:
        from openpyxl import load_workbook
    except Exception as exc:  # pragma: no cover - optional dependency
        raise RuntimeError(
            "Reading .xlsx files requires the 'openpyxl' package."
        ) from exc

    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    rows_iter = sheet.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration:
        workbook.close()
        return [], []
    headers = [str(cell).strip() for cell in header_row if cell is not None]
    rows: list[dict[str, str]] = []
    for values in rows_iter:
        row: dict[str, str] = {}
        for index, value in enumerate(values):
            if index >= len(headers):
                break
            if value is None:
                row[headers[index]] = ""
            elif isinstance(value, float) and value.is_integer():
                row[headers[index]] = str(int(value))
            else:
                row[headers[index]] = str(value).strip()
        if any(row.values()):
            rows.append(row)
    workbook.close()
    return headers, rows


def _split_list(value: str) -> list[str]:
    parts = [p.strip() for p in value.split(LIST_SEPARATOR)]
    return [p for p in parts if p]


def _first_present(row: dict[str, str], candidates: Iterable[str]) -> str:
    lower = {k.lower(): v for k, v in row.items()}
    for candidate in candidates:
        value = lower.get(candidate, "")
        if value:
            return value
    return ""


def parse_upload_rows(rows: Sequence[dict[str, str]]) -> list[dict]:
    """Turn spreadsheet rows into upload batches.

    Each returned item has an ``identifier``, a list of local ``files`` and a
    ``metadata`` dict.  Multi-value metadata may be separated with ``|``.
    """
    batches: list[dict] = []
    for index, row in enumerate(rows, start=1):
        identifier = _first_present(row, IDENTIFIER_COLUMNS)
        if not identifier:
            log.warning("Skipping spreadsheet row %d: no identifier column.", index)
            continue
        files = _split_list(_first_present(row, FILE_COLUMNS))
        metadata: dict = {}
        for key, value in row.items():
            if not key or key.lower() in RESERVED_COLUMNS or not value:
                continue
            parts = _split_list(value) if LIST_SEPARATOR in value else [value]
            metadata[key] = parts if len(parts) > 1 else parts[0]
        batches.append({"identifier": identifier, "files": files, "metadata": metadata})
    return batches


def parse_metadata_rows(rows: Sequence[dict[str, str]]) -> list[tuple[str, dict]]:
    """Turn spreadsheet rows into ``(identifier, metadata)`` pairs."""
    updates: list[tuple[str, dict]] = []
    for index, row in enumerate(rows, start=1):
        identifier = _first_present(row, IDENTIFIER_COLUMNS)
        if not identifier:
            log.warning("Skipping spreadsheet row %d: no identifier column.", index)
            continue
        metadata: dict = {}
        for key, value in row.items():
            if not key or key.lower() in RESERVED_COLUMNS or not value:
                continue
            parts = _split_list(value) if LIST_SEPARATOR in value else [value]
            metadata[key] = parts if len(parts) > 1 else parts[0]
        if metadata:
            updates.append((identifier, metadata))
    return updates


def export_results(path: str, rows: Sequence[dict]) -> str:
    """Write search results to CSV or JSON based on the file extension."""
    file_path = Path(path)
    suffix = file_path.suffix.lower()
    if suffix == ".json":
        file_path.write_text(json.dumps(list(rows), indent=2, default=str), encoding="utf-8")
        return str(file_path)

    fields: list[str] = []
    for row in rows:
        for key in row:
            if key not in fields:
                fields.append(key)
    with file_path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        for row in rows:
            writer.writerow({key: _flatten(row.get(key, "")) for key in fields})
    return str(file_path)


def _flatten(value) -> str:
    if isinstance(value, (list, tuple)):
        return LIST_SEPARATOR.join(str(v) for v in value)
    if isinstance(value, dict):
        return json.dumps(value, default=str)
    return "" if value is None else str(value)


def write_template(path: str, headers: Sequence[str], sample: Sequence[dict] | None = None) -> str:
    """Write a starter CSV template for batch operations."""
    with Path(path).open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(headers))
        writer.writeheader()
        for row in sample or []:
            writer.writerow(row)
    return path


__all__ = [
    "read_table",
    "parse_upload_rows",
    "parse_metadata_rows",
    "export_results",
    "write_template",
    "IDENTIFIER_COLUMNS",
    "FILE_COLUMNS",
]
